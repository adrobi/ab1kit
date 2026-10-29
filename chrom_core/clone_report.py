"""CSV and Excel reports for clone verification (shared by the UI and the CLI)."""

from __future__ import annotations

from chrom_core.batch import csv_field
from chrom_core.clones import CloneDifference, CloneReport

# verdict -> (light, dark) hex colours; the single source of the status palette.
STATUS_HEX: dict[str, tuple[str, str]] = {
    "ok": ("#c6efce", "#2e7d32"),
    "errors": ("#ffc7ce", "#c62828"),
    "incomplete": ("#ffeb9c", "#f9a825"),
    "mixed": ("#e1d5f0", "#7b1fa2"),
    "no_data": ("#e0e0e0", "#757575"),
}

CLONES_HEADER = ("clone", "verdict", "region_coverage_pct", "both_strands_pct",
                 "differences", "reads")
DIFFERENCES_HEADER = ("clone", "position", "kind", "ref", "read", "effect",
                      "support", "against", "status", "note")


def _plural(n: int, word: str) -> str:
    return f"{n} {word}" if n == 1 else f"{n} {word}s"


def _note(d: CloneDifference) -> str:
    """Short remark: ``disputed: N read(s) against`` / ``low quality: N read(s)``,
    plus ``outside region`` (joined by `` / ``); empty for a confirmed in-region one."""
    parts = []
    if d.status == "disputed":
        parts.append(f"disputed: {_plural(d.against, 'read')} against")
    elif d.status == "low_quality":
        parts.append(f"low quality: {_plural(d.low_quality, 'read')}")
    if not d.in_region:
        parts.append("outside region")
    return " / ".join(parts)


def _clone_row(r: CloneReport) -> tuple:
    # ``reads``: file names of every read of the clone (used or not), joined by "|".
    names = "|".join(read.ref.path.name for read in r.reads)
    return (r.clone, r.verdict, round(r.region_coverage_pct, 1), round(r.both_strands_pct, 1),
            len(r.differences), names)


def _difference_rows(reports: list[CloneReport]):
    for r in reports:
        for d in r.differences:
            # position is 1-based; an insertion is reported before this position
            yield (r.clone, d.ref_pos + 1, d.kind, d.ref_base, d.read_base, d.effect,
                   d.support, d.against, d.status, _note(d))


def _fmt(value) -> str:
    return f"{value:.1f}" if isinstance(value, float) else str(value)


def _csv(header, rows) -> str:
    lines = [",".join(header)]
    for row in rows:
        lines.append(",".join(csv_field(_fmt(v), sep=",") for v in row))
    return "\n".join(lines) + "\n"


def clones_csv(reports: list[CloneReport]) -> str:
    """One row per clone, comma-separated."""
    return _csv(CLONES_HEADER, (_clone_row(r) for r in reports))


def differences_csv(reports: list[CloneReport]) -> str:
    """One row per difference of every clone (outside-region ones included)."""
    return _csv(DIFFERENCES_HEADER, _difference_rows(reports))


def clone_report_xlsx(path, reports: list[CloneReport], construct_name: str) -> None:
    """Excel report: sheets ``Clones`` and ``Differences``; the verdict cell is filled
    with its light ``STATUS_HEX`` colour. ``construct_name`` is the workbook title."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise RuntimeError("openpyxl is required for Excel reports: pip install openpyxl") from exc

    book = Workbook()
    book.properties.title = construct_name
    clones = book.active
    clones.title = "Clones"
    diffs = book.create_sheet("Differences")
    for sheet, header in ((clones, CLONES_HEADER), (diffs, DIFFERENCES_HEADER)):
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    for r in reports:
        clones.append(_clone_row(r))
        palette = STATUS_HEX.get(r.verdict)
        if palette:
            colour = palette[0].lstrip("#").upper()
            clones.cell(row=clones.max_row, column=2).fill = PatternFill(
                "solid", start_color=colour, end_color=colour)
    for row in _difference_rows(reports):
        diffs.append(row)
    for sheet, widths in ((clones, (16, 12, 20, 18, 12, 50)),
                          (diffs, (16, 10, 14, 8, 8, 14, 9, 9, 12, 30))):
        for i, width in enumerate(widths):
            sheet.column_dimensions[chr(ord("A") + i)].width = width
    book.save(path)
