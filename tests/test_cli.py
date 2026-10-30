"""Command-line subcommands: verify-clones, qc, export (thin shell over the core)."""

import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from Bio import SeqIO
from openpyxl import load_workbook

import chrom_cli
from clone_fixtures import COI_GB, DATA, coi_reference, mutate

ROOT = Path(__file__).resolve().parent.parent
READS = ["Achl_ACHLO006-09_1_F.ab1", "Achl_ACHLO006-09_2_R.ab1",
         "Achl_ACHLO007-09_1_F.ab1", "Achl_ACHLO007-09_2_R.ab1"]


@pytest.fixture
def reads_dir(tmp_path: Path) -> Path:
    d = tmp_path / "Данные клоны"
    d.mkdir()
    for name in READS:
        shutil.copy(DATA / "real" / name, d / name)
    return d


def test_verify_clones_xlsx(reads_dir: Path, tmp_path: Path, capsys) -> None:
    out = tmp_path / "отчёт.xlsx"
    code = chrom_cli.main(["verify-clones", "--construct", str(COI_GB), str(reads_dir),
                           "-o", str(out)])
    assert code == 0
    book = load_workbook(out)
    assert book.sheetnames == ["Clones", "Differences"]
    shown = capsys.readouterr().out
    assert "Achl_ACHLO006-09" in shown and "ok" in shown


def test_verify_clones_csv_region_and_errors(reads_dir: Path, tmp_path: Path, capsys) -> None:
    ref = tmp_path / "mutated.fasta"
    seq = coi_reference()
    ref.write_text(">mut\n" + mutate(seq, 250, "sub", "A" if seq[250] != "A" else "C") + "\n")
    out = tmp_path / "rep.csv"
    code = chrom_cli.main(["verify-clones", "--construct", str(ref), "--region", "7-647",
                           "--linear", "--min-q", "20", "--het-ratio", "0.4",
                           str(reads_dir), "-o", str(out)])
    assert code == 1  # every clone carries the planted difference
    assert "errors" in out.read_text(encoding="utf-8")
    diffs = tmp_path / "rep_differences.csv"
    assert diffs.exists() and "251" in diffs.read_text(encoding="utf-8")
    assert "errors" in capsys.readouterr().out


def test_verify_clones_input_errors(reads_dir: Path, tmp_path: Path, capsys) -> None:
    empty = tmp_path / "empty"
    empty.mkdir()
    out = tmp_path / "x.xlsx"
    base = ["verify-clones", "-o", str(out)]
    assert chrom_cli.main(base + ["--construct", str(COI_GB), str(empty)]) == 2
    bad = tmp_path / "bad.gb"
    bad.write_text("not a sequence file")
    assert chrom_cli.main(base + ["--construct", str(bad), str(reads_dir)]) == 2
    assert chrom_cli.main(base + ["--construct", str(COI_GB), "--region", "9-3",
                                  "--linear", str(reads_dir)]) == 2
    assert chrom_cli.main(base + ["--construct", str(COI_GB), "--region", "abc",
                                  str(reads_dir)]) == 2
    assert chrom_cli.main(base + ["--construct", str(COI_GB), "--pattern", "(",
                                  str(reads_dir)]) == 2
    assert chrom_cli.main(base + ["--construct", str(COI_GB),
                                  str(tmp_path / "missing")]) == 2
    assert capsys.readouterr().err.count("ab1kit:") >= 6
    assert not out.exists()


def test_unassigned_files_reported_on_stderr(reads_dir: Path, tmp_path: Path, capsys) -> None:
    shutil.copy(DATA / "real" / READS[0], reads_dir / "!.ab1")
    out = tmp_path / "r.csv"
    chrom_cli.main(["verify-clones", "--construct", str(COI_GB), str(reads_dir), "-o", str(out)])
    assert "!.ab1" in capsys.readouterr().err


def test_qc_stdout_and_file(reads_dir: Path, tmp_path: Path, capsys) -> None:
    assert chrom_cli.main(["qc", str(reads_dir)]) == 0
    printed = capsys.readouterr().out
    assert printed.startswith("file;length;mean_q") and READS[0] in printed
    out = tmp_path / "qc.csv"
    assert chrom_cli.main(["qc", str(reads_dir), "-o", str(out)]) == 0
    assert READS[1] in out.read_text(encoding="utf-8")


def test_export(reads_dir: Path, tmp_path: Path) -> None:
    out = tmp_path / "вывод"
    assert chrom_cli.main(["export", str(reads_dir), "--trim", "--fasta", "-o", str(out)]) == 0
    files = sorted(p.name for p in out.glob("*.trimmed.fasta"))
    assert len(files) == len(READS)
    out2 = tmp_path / "fq"
    assert chrom_cli.main(["export", str(reads_dir), "--trim", "-o", str(out2)]) == 0
    assert len(list(out2.glob("*.trimmed.fastq"))) == len(READS)


@pytest.mark.parametrize("argv", [[], ["x.ab1"]])
def test_gui_path_unchanged(monkeypatch, argv) -> None:
    import ui.main_window as mw
    seen = []
    monkeypatch.setattr(mw, "run", lambda a: seen.append(a) or 0)
    assert chrom_cli.main(argv) == 0
    assert seen == [argv]


def _run_cli(*args, **kw):
    return subprocess.run([sys.executable, "-m", "chrom_cli", *map(str, args)],
                          capture_output=True, cwd=ROOT, **kw)


def test_cli_does_not_import_qt_and_survives_legacy_console(reads_dir: Path, tmp_path: Path) -> None:
    out = tmp_path / "r.csv"
    code = (
        "import sys, chrom_cli\n"
        f"rc = chrom_cli.main(['verify-clones','--construct',r'{COI_GB}',r'{reads_dir}','-o',r'{out}'])\n"
        "assert not any(m == 'PySide6' or m.startswith('PySide6.') for m in sys.modules), 'Qt imported'\n"
        "sys.exit(rc)\n"
    )
    env = {**__import__("os").environ, "PYTHONIOENCODING": "cp1251"}
    env.pop("PYTHONUTF8", None)
    proc = subprocess.run([sys.executable, "-c", code], capture_output=True, cwd=ROOT, env=env)
    assert proc.returncode == 0, proc.stderr.decode("utf-8", "replace")


def test_cyrillic_paths_on_legacy_console(tmp_path: Path) -> None:
    d = tmp_path / "Данные клоны"
    d.mkdir()
    for name in READS[:2]:
        shutil.copy(DATA / "real" / name, d / name)
    shutil.copy(DATA / "real" / READS[0], d / "Клон ж_1_F.ab1")
    shutil.copy(DATA / "real" / READS[0], d / "Ё!.ab1")
    env = {**__import__("os").environ, "PYTHONIOENCODING": "cp437"}  # cannot encode Cyrillic
    env.pop("PYTHONUTF8", None)
    proc = subprocess.run(
        [sys.executable, "-m", "chrom_cli", "verify-clones", "--construct", str(COI_GB),
         str(d), "-o", str(tmp_path / "o.csv")],
        capture_output=True, cwd=ROOT, env=env)
    assert proc.returncode in (0, 1), proc.stderr.decode("utf-8", "replace")
    assert b"Traceback" not in proc.stderr
    assert "Клон ж".encode() in proc.stdout  # utf-8 forced regardless of the console
    assert "Ё!.ab1".encode() in proc.stderr
