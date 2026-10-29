"""CSV / Excel reports for clone verification, built from real reads (HM417896 COI)."""

from __future__ import annotations

import csv
import io
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from chrom_core import batch
from chrom_core.clone_report import (
    STATUS_HEX,
    clone_report_xlsx,
    clones_csv,
    differences_csv,
)
from chrom_core.clones import Grouping, ReadRef, verify_all, verify_clone
from chrom_core.reference import Construct, load_construct
from clone_fixtures import COI_GB, DATA, REAL_READS, coi_reference, mutate

F = DATA / "real" / "Achl_ACHLO006-09_1_F.ab1"
R = DATA / "real" / "Achl_ACHLO006-09_2_R.ab1"
REGION = (6, 647)
P = 250        # inside the region
OUTSIDE = 652  # covered by F only, outside the region

assert F in REAL_READS and R in REAL_READS


@pytest.fixture(scope="module")
def coi() -> str:
    return coi_reference()


def _reads(clone: str, *paths_primers) -> list[ReadRef]:
    return [ReadRef(p, clone, primer) for p, primer in paths_primers]


@pytest.fixture(scope="module")
def reports(coi):
    """c1: clean (ok) and c2: substitution at P and outside-region one at OUTSIDE."""
    clean = verify_clone(Construct("c", coi, False), REGION, _reads("c1", (F, "F"), (R, "R")))
    bad_ref = mutate(mutate(coi, P, "sub", "G" if coi[P] != "G" else "A"),
                     OUTSIDE, "sub", "G" if coi[OUTSIDE] != "G" else "A")
    bad = verify_clone(Construct("c", bad_ref, False), REGION, _reads("c2", (F, "F"), (R, "R")))
    return [clean, bad]


def _rows(text: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(text)))


def test_clones_csv_header_and_rows(reports):
    rows = _rows(clones_csv(reports))
    assert rows[0] == ["clone", "verdict", "region_coverage_pct", "both_strands_pct",
                       "differences", "reads"]
    assert len(rows) == 3
    c1, c2 = rows[1], rows[2]
    assert c1[:2] == ["c1", "ok"]
    assert c1[2] == "100.0" and c1[3] == "97.0"   # F and R overlap on most of the region
    assert c1[4] == "0"
    assert c1[5] == f"{F.name}|{R.name}"
    assert c2[:2] == ["c2", "errors"]
    assert c2[4] == str(len(reports[1].differences))


def test_clones_csv_empty():
    assert clones_csv([]).splitlines() == [
        "clone,verdict,region_coverage_pct,both_strands_pct,differences,reads"]


def test_differences_csv_header_position_and_notes(reports):
    rows = _rows(differences_csv(reports))
    assert rows[0] == ["clone", "position", "kind", "ref", "read", "effect",
                       "support", "against", "status", "note"]
    body = rows[1:]
    assert len(body) == len(reports[1].differences) == 2   # clean clone contributes none
    by_pos = {int(r[1]): r for r in body}
    assert set(by_pos) == {P + 1, OUTSIDE + 1}              # 1-based
    inside, outside = by_pos[P + 1], by_pos[OUTSIDE + 1]
    assert inside[0] == "c2" and inside[2] == "substitution"
    assert inside[8] == "confirmed" and inside[9] == ""
    assert inside[6] == "2" and inside[7] == "0"
    assert outside[9] == "outside region"


def test_differences_note_variants():
    from chrom_core.clone_report import _note
    from chrom_core.clones import CloneDifference

    def diff(**kw):
        base = dict(ref_pos=1, kind="substitution", ref_base="A", read_base="G", support=1,
                    against=0, low_quality=0, status="confirmed", in_region=True)
        return CloneDifference(**{**base, **kw})

    assert _note(diff()) == ""
    assert _note(diff(status="disputed", against=1)) == "disputed: 1 read against"
    assert _note(diff(status="disputed", against=3)) == "disputed: 3 reads against"
    assert _note(diff(status="low_quality", support=0, low_quality=2)) == "low quality: 2 reads"
    assert _note(diff(status="low_quality", support=0, low_quality=1, in_region=False)) == \
        "low quality: 1 read / outside region"


def test_csv_quotes_commas_and_quotes():
    from chrom_core.clones import CloneReport
    weird = CloneReport('a,"b"', (), {}, (), (), "no_data")
    rows = _rows(clones_csv([weird]))
    assert rows[1][0] == 'a,"b"'
    assert len(rows[1]) == 6
    assert batch.csv_field("x,y", sep=",") == '"x,y"'
    assert batch.csv_field("x,y") == "x,y"


def test_xlsx_sheets_and_fill(reports, tmp_path):
    path = tmp_path / "r.xlsx"
    clone_report_xlsx(path, reports, "COI")
    book = load_workbook(path)
    assert book.sheetnames == ["Clones", "Differences"]
    clones, diffs = book["Clones"], book["Differences"]
    assert [c.value for c in clones[1]] == ["clone", "verdict", "region_coverage_pct",
                                            "both_strands_pct", "differences", "reads"]
    assert clones.max_row == 3
    assert clones["B2"].value == "ok" and clones["B3"].value == "errors"
    assert clones["B3"].fill.fgColor.rgb[-6:] == STATUS_HEX["errors"][0].lstrip("#").upper()
    assert clones["B2"].fill.fgColor.rgb[-6:] == STATUS_HEX["ok"][0].lstrip("#").upper()
    assert clones["B1"].font.bold
    assert clones["C2"].value == 100.0
    assert diffs.max_row == 3
    assert diffs["B2"].value in (P + 1, OUTSIDE + 1)
    assert book.properties.title == "COI"


def test_status_hex_palette():
    assert set(STATUS_HEX) == {"ok", "errors", "incomplete", "mixed", "no_data"}
    for light, dark in STATUS_HEX.values():
        for colour in (light, dark):
            assert len(colour) == 7 and colour[0] == "#"
            int(colour[1:], 16)
    assert len({light for light, _ in STATUS_HEX.values()}) == 5


def test_end_to_end_effect_from_cds_construct(coi, tmp_path):
    construct = load_construct(COI_GB)
    assert any(f.kind == "CDS" for f in construct.features)
    alt = "G" if coi[P] != "G" else "A"
    mutated = replace(construct, seq=mutate(construct.seq, P, "sub", alt))
    grouping = Grouping([ReadRef(F, "c1", "F"), ReadRef(R, "c1", "R")])

    [report] = verify_all(mutated, REGION, grouping)
    [diff] = [d for d in report.differences if d.in_region]
    assert diff.effect

    [plain] = verify_all(mutated, REGION, grouping, annotate=None)
    assert [d.effect for d in plain.differences] == [""] * len(plain.differences)

    rows = _rows(differences_csv([report]))
    row = next(r for r in rows[1:] if int(r[1]) == P + 1)
    assert row[5] == diff.effect
